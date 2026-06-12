import openpyxl

wb = openpyxl.load_workbook('data/Metric Evaluation Text to Image.xlsx')
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f"Sheet: {sheet_name}")
    print(f"  Prompt (Row 2 Col 3): {sheet.cell(row=2, column=3).value[:60]}...")
    print(f"  Model (Row 2 Col 2): {sheet.cell(row=2, column=2).value}")
